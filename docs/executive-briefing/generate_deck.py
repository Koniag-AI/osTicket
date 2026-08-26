#!/usr/bin/env python3
"""Executive deck: AI-governed service delivery (ServiceNow/Jira -> Bedrock -> Devin -> Harness)."""

from pptx import Presentation
from pptx.util import Inches as I, Pt
from pptx.dml.color import RGBColor as C
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.dml import MSO_LINE_DASH_STYLE

# ---------------- design system ----------------
BG1, BG2 = C(0x08, 0x0E, 0x1C), C(0x10, 0x1B, 0x38)
PANEL = C(0x14, 0x1F, 0x3A)
PANEL2 = C(0x0F, 0x18, 0x2E)
WHITE = C(0xFF, 0xFF, 0xFF)
MUTED = C(0x9A, 0xAC, 0xC8)
DIM = C(0x6B, 0x7C, 0x99)
CYAN = C(0x22, 0xD3, 0xEE)
BLUE = C(0x3B, 0x82, 0xF6)
VIOLET = C(0x8B, 0x5C, 0xF6)
TEAL = C(0x2D, 0xD4, 0xBF)
AMBER = C(0xF5, 0x9E, 0x0B)
GREEN = C(0x34, 0xD3, 0x99)
FONT = "Segoe UI"

STAGES = [
    ("1", "INTAKE", "ServiceNow  ·  Jira", BLUE,
     ["One front door for every request",
      "Normalized to a common schema",
      "\u201cReady for AI\u201d fires the pipeline"]),
    ("2", "QUALIFY", "Amazon Bedrock", VIOLET,
     ["Links repeats to the open parent",
      "Scores completeness and detail",
      "Rates risk and AI eligibility"]),
    ("3", "RESOLVE", "Devin", TEAL,
     ["Automations start work instantly",
      "Playbooks make fixes repeatable",
      "Ships a tested pull request"]),
    ("4", "PROVE", "Harness", AMBER,
     ["SAST, dependency, IaC policy scans",
      "Named human approves release",
      "Every promotion logged for audit"]),
    ("5", "COMPOUND", "Back to the ticket", GREEN,
     ["Resolution written to the ticket",
      "Fixes become playbooks",
      "Autonomy measured per stage"]),
]

prs = Presentation()
prs.slide_width, prs.slide_height = I(13.333), I(7.5)
BLANK = prs.slide_layouts[6]
SW = 13.333


# ---------------- primitives ----------------
def sl():
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, I(13.333), I(7.5))
    r.line.fill.background()
    r.shadow.inherit = False
    f = r.fill
    f.gradient()
    f.gradient_angle = 45.0
    f.gradient_stops[0].color.rgb = BG1
    f.gradient_stops[1].color.rgb = BG2
    return s


def box(s, x, y, w, h, fill=None, line=None, radius=0.05, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
        lw=1.0, grad=None, angle=0.0):
    sh = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    sh.shadow.inherit = False
    if grad:
        f = sh.fill
        f.gradient()
        f.gradient_angle = angle
        f.gradient_stops[0].color.rgb = grad[0]
        f.gradient_stops[1].color.rgb = grad[1]
    elif fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(lw)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.text_frame.word_wrap = True
    return sh


def txt(s, x, y, w, h, runs, size=12, color=WHITE, bold=False, align=PP_ALIGN.LEFT,
        anchor=MSO_ANCHOR.TOP, space=6, line=1.15, caps=False, spacing=None):
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    items = runs if isinstance(runs, list) else [runs]
    for i, t in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line
        r = p.add_run()
        r.text = t.upper() if caps else t
        fnt = r.font
        fnt.name = FONT
        fnt.size = Pt(size)
        fnt.bold = bold
        fnt.color.rgb = color
        if spacing:
            from pptx.oxml.ns import qn
            fnt._rPr.set("spc", str(int(spacing * 100)))
    return tb


def header(s, kicker, title, sub=None, accent=CYAN):
    box(s, 0.55, 0.52, 0.055, 0.30, fill=accent, radius=0.5)
    txt(s, 0.75, 0.53, 8.0, 0.3, kicker, size=10.5, color=accent, bold=True, caps=True, spacing=1.6)
    txt(s, 0.55, 0.86, 11.2, 0.6, title, size=27, color=WHITE, bold=True, space=0)
    if sub:
        txt(s, 0.55, 1.42, 11.6, 0.35, sub, size=13, color=MUTED, space=0)


def chip(s, x, y, w, text, color, h=0.34, size=10, fill=None, bold=True):
    sh = box(s, x, y, w, h, fill=fill if fill else PANEL2, line=color, radius=0.5, lw=0.9)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = I(0.08)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return sh


def footer(s, label):
    box(s, 0.55, 7.05, 12.23, 0.012, fill=C(0x1E, 0x2C, 0x4C))
    txt(s, 0.55, 7.12, 8.0, 0.25, label, size=9, color=DIM, space=0)
    txt(s, 9.0, 7.12, 3.78, 0.25, "Human in the loop at every gate", size=9,
        color=DIM, align=PP_ALIGN.RIGHT, space=0)


def arrow(s, x, y, color=C(0x2A, 0x3B, 0x5E), w=0.22, h=0.22):
    sh = box(s, x, y, w, h, fill=color, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
    sh.rotation = 90
    return sh


# ================= 1. title =================
s = sl()
box(s, 0, 0, 13.333, 0.10, grad=(VIOLET, CYAN), shape=MSO_SHAPE.RECTANGLE, angle=0)
box(s, 8.9, 1.1, 4.0, 4.0, fill=C(0x14, 0x22, 0x44), radius=0.5)
box(s, 9.35, 1.55, 3.1, 3.1, fill=C(0x18, 0x2A, 0x54), radius=0.5)
box(s, 9.9, 2.1, 2.0, 2.0, grad=(VIOLET, CYAN), radius=0.5, angle=45)
txt(s, 9.9, 2.72, 2.0, 0.8, "AI", size=48, color=WHITE, bold=True, align=PP_ALIGN.CENTER, space=0)

txt(s, 0.85, 1.75, 8.0, 0.4, "Executive briefing", size=12, color=CYAN, bold=True, caps=True, spacing=2.2)
txt(s, 0.85, 2.25, 8.0, 2.0,
    ["AI-Governed", "Service Delivery"], size=46, color=WHITE, bold=True, space=2, line=1.02)
box(s, 0.9, 4.35, 1.5, 0.045, fill=VIOLET)
txt(s, 0.85, 4.62, 7.6, 1.0,
    "From ticket intake to compliant release \u2014 autonomous where it is safe, "
    "human-approved where it counts.", size=15.5, color=MUTED, space=0, line=1.25)
cx = 0.85
for t, col, w in [("ServiceNow", BLUE, 1.35), ("Jira", BLUE, 0.95), ("Amazon Bedrock", VIOLET, 1.75),
                  ("Devin", TEAL, 1.10), ("Harness", AMBER, 1.25)]:
    chip(s, cx, 5.75, w, t, col, size=9.5)
    cx += w + 0.14
footer(s, "Koniag \u00b7 AI service delivery architecture")

# ================= 2. hero flow =================
s = sl()
header(s, "End-to-end flow", "Five stages, one governed path.",
       "Bedrock qualifies the work \u00b7 Devin does the work \u00b7 Harness proves it is safe \u00b7 a person approves the release")

CX, CW, GAPX = 0.55, 2.36, 0.135
CY, CH = 1.95, 3.05
for i, (num, title, tool, col, bullets) in enumerate(STAGES):
    x = CX + i * (CW + GAPX)
    box(s, x, CY, CW, CH, grad=(PANEL, PANEL2), radius=0.06, angle=90, line=C(0x24, 0x33, 0x55), lw=0.75)
    box(s, x, CY, CW, 0.055, fill=col, shape=MSO_SHAPE.RECTANGLE)
    box(s, x + 0.20, CY + 0.28, 0.34, 0.34, fill=col, radius=0.5)
    txt(s, x + 0.20, CY + 0.345, 0.34, 0.25, num, size=12, color=BG1, bold=True, align=PP_ALIGN.CENTER, space=0)
    txt(s, x + 0.62, CY + 0.31, CW - 0.8, 0.3, title, size=15.5, color=WHITE, bold=True, space=0, spacing=1.0)
    txt(s, x + 0.20, CY + 0.74, CW - 0.4, 0.26, tool, size=10.5, color=col, bold=True, space=0)
    box(s, x + 0.20, CY + 1.06, CW - 0.4, 0.01, fill=C(0x25, 0x35, 0x58))
    ty = CY + 1.28
    for b in bullets:
        box(s, x + 0.20, ty + 0.075, 0.055, 0.055, fill=col, radius=0.5)
        txt(s, x + 0.36, ty, CW - 0.52, 0.55, b, size=10, color=MUTED, space=0, line=1.2)
        ty += 0.58
    if i < 4:
        arrow(s, x + CW + 0.005, CY + 0.36, color=col)

# return paths
txt(s, 0.55, 5.28, 7.0, 0.25, "Return paths that protect the pipeline", size=10.5,
    color=CYAN, bold=True, caps=True, spacing=1.4, space=0)
loops = [
    ("\u21ba  Duplicate or thin ticket \u2192 returned at the gate", VIOLET),
    ("\u21ba  Failed scan or review \u2192 Devin reworks the fix", AMBER),
    ("\u21ba  Resolved outcomes \u2192 sharper playbooks", GREEN),
]
lw = 4.02
for i, (t, col) in enumerate(loops):
    sh = box(s, 0.55 + i * (lw + 0.10), 5.60, lw, 0.44, fill=PANEL2, line=col, radius=0.35, lw=0.9)
    tf = sh.text_frame
    tf.margin_left = I(0.14)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = t
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(9.5), True, col

# governance band
box(s, 0.55, 6.20, 12.23, 0.70, grad=(C(0x16, 0x24, 0x46), C(0x0E, 0x18, 0x30)), radius=0.10, angle=0)
txt(s, 0.75, 6.32, 2.6, 0.5, ["HUMAN CONTROL", "POINTS"], size=10, color=CYAN, bold=True,
    spacing=1.4, space=0, line=1.15)
ctrl = [("Intake owner", "business priority"), ("Triage thresholds", "tuned, not guessed"),
        ("Pull request review", "every change is read"), ("Harness approval", "release needs a person"),
        ("Audit trail", "ticket \u2192 PR \u2192 deploy")]
cw = 9.4 / 5
for i, (h, sub) in enumerate(ctrl):
    x = 3.30 + i * cw
    box(s, x, 6.44, 0.08, 0.08, fill=CYAN, radius=0.5)
    txt(s, x + 0.17, 6.36, cw - 0.25, 0.25, h, size=10.5, color=WHITE, bold=True, space=0)
    txt(s, x + 0.17, 6.60, cw - 0.25, 0.25, sub, size=9, color=DIM, space=0)
footer(s, "Architecture at a glance")

# ================= 3. the gate =================
s = sl()
header(s, "Stage 2 \u00b7 the gate", "No work starts until the ticket earns it.",
       "Amazon Bedrock evaluates every ticket before a single engineering cycle \u2014 human or AI \u2014 is spent.",
       accent=VIOLET)
tests = [
    ("Is it new?", "Semantic comparison against ticket history clusters repeats and links them to the open parent.", VIOLET),
    ("Is it complete?", "Repro steps, logs, environment, affected system and scope are scored against a required-detail standard.", VIOLET),
    ("Is it eligible?", "Category, blast radius and risk decide autonomous fix, assisted fix, or human-only.", VIOLET),
]
for i, (h, b, col) in enumerate(tests):
    y = 2.00 + i * 1.15
    box(s, 0.55, y, 7.05, 1.00, grad=(PANEL, PANEL2), radius=0.10, angle=90, line=C(0x28, 0x35, 0x5C), lw=0.75)
    box(s, 0.55, y, 0.05, 1.00, fill=col, shape=MSO_SHAPE.RECTANGLE)
    txt(s, 0.85, y + 0.18, 0.5, 0.4, f"0{i+1}", size=20, color=col, bold=True, space=0)
    txt(s, 1.55, y + 0.17, 5.85, 0.3, h, size=15, color=WHITE, bold=True, space=0)
    txt(s, 1.55, y + 0.50, 5.85, 0.45, b, size=10.5, color=MUTED, space=0, line=1.2)

box(s, 7.95, 2.00, 4.83, 3.45, grad=(C(0x17, 0x25, 0x48), C(0x0E, 0x18, 0x30)), radius=0.07, angle=90,
    line=C(0x28, 0x35, 0x5C), lw=0.75)
txt(s, 8.25, 2.22, 4.2, 0.3, "Gate outcomes", size=13, color=WHITE, bold=True, space=0)
out = [("Returned", "duplicate, or too thin to act on", AMBER),
       ("Escalated", "high risk or policy-restricted \u2192 human-only", BLUE),
       ("Qualified", "scoped, evidenced, safe \u2192 handed to Devin", GREEN)]
for i, (h, b, col) in enumerate(out):
    y = 2.68 + i * 0.88
    box(s, 8.25, y, 4.23, 0.72, fill=PANEL2, radius=0.12, line=col, lw=0.9)
    box(s, 8.42, y + 0.14, 0.08, 0.44, fill=col, radius=0.5)
    txt(s, 8.62, y + 0.11, 3.7, 0.25, h, size=11.5, color=col, bold=True, space=0)
    txt(s, 8.62, y + 0.36, 3.7, 0.3, b, size=9.5, color=MUTED, space=0)

box(s, 0.55, 5.72, 12.23, 1.06, grad=(C(0x1B, 0x14, 0x3C), C(0x0E, 0x18, 0x30)), radius=0.10, angle=0,
    line=VIOLET, lw=0.9)
txt(s, 0.85, 5.90, 11.6, 0.35, "Why this matters", size=10, color=VIOLET, bold=True, caps=True, spacing=1.4, space=2)
txt(s, 0.85, 6.20, 11.6, 0.5,
    "Repeat and low-quality tickets are the largest hidden cost in service desks. Screening them at the gate "
    "protects engineering capacity, keeps automation trustworthy, and gives the business a clean signal on real demand.",
    size=12, color=MUTED, space=0, line=1.2)
footer(s, "Qualification \u00b7 Amazon Bedrock")

# ================= 4. the work =================
s = sl()
header(s, "Stage 3 \u00b7 the work", "Devin executes the qualified ticket, end to end.",
       "Automations decide when work starts. Playbooks decide how it is done. Every run ships evidence, not just code.",
       accent=TEAL)
steps = [("Trigger", "Qualified ticket starts a session automatically"),
         ("Playbook", "The tested procedure for that work type is applied"),
         ("Change", "Code, config or IaC updated in the target repo"),
         ("Verify", "Tests, lint and type checks run before hand-off"),
         ("Evidence", "Pull request with reasoning, diff and test output"),
         ("Write-back", "Ticket updated with status and release link")]
bw, bg_ = 1.90, 0.115
for i, (h, b) in enumerate(steps):
    x = 0.55 + i * (bw + bg_)
    box(s, x, 2.02, bw, 1.62, grad=(PANEL, PANEL2), radius=0.10, angle=90, line=C(0x24, 0x33, 0x55), lw=0.75)
    box(s, x + 0.22, 2.24, 0.30, 0.30, fill=TEAL, radius=0.5)
    txt(s, x + 0.22, 2.30, 0.30, 0.22, str(i + 1), size=11, color=BG1, bold=True, align=PP_ALIGN.CENTER, space=0)
    txt(s, x + 0.22, 2.68, bw - 0.44, 0.28, h, size=12.5, color=WHITE, bold=True, space=0)
    txt(s, x + 0.22, 2.98, bw - 0.44, 0.6, b, size=9.5, color=MUTED, space=0, line=1.2)
    if i < 5:
        arrow(s, x + bw - 0.02, 2.30, color=TEAL, w=0.16, h=0.16)

pairs = [("Automations", "Event-driven start", TEAL,
          ["Fires on ticket creation or a \u201cReady for AI\u201d transition",
           "No queue, no triage meeting, no assignment delay",
           "Scope and guardrails come from the ticket itself"]),
         ("Playbooks", "Repeatable method", CYAN,
          ["Curated, tested procedure per class of work",
           "Same steps, same checks, every single time",
           "Improved as outcomes come back from stage 5"])]
for i, (h, tag, col, bl) in enumerate(pairs):
    x = 0.55 + i * 6.23
    box(s, x, 3.95, 6.00, 2.10, grad=(PANEL, PANEL2), radius=0.07, angle=90, line=C(0x24, 0x33, 0x55), lw=0.75)
    txt(s, x + 0.30, 4.18, 3.4, 0.3, h, size=15.5, color=WHITE, bold=True, space=0)
    chip(s, x + 4.05, 4.16, 1.65, tag, col, h=0.30, size=9)
    yy = 4.66
    for b in bl:
        box(s, x + 0.32, yy + 0.08, 0.07, 0.07, fill=col, radius=0.5)
        txt(s, x + 0.52, yy, 5.2, 0.3, b, size=11, color=MUTED, space=0)
        yy += 0.42
txt(s, 0.55, 6.30, 12.23, 0.5,
    "Devin never releases anything on its own \u2014 stage 3 always ends at a reviewable pull request.",
    size=12, color=WHITE, bold=True, space=0)
footer(s, "Autonomous execution \u00b7 Devin")

# ================= 5. the proof =================
s = sl()
header(s, "Stage 4 \u00b7 the proof", "Compliance enforced by the pipeline.",
       "Harness runs the same governed gate on AI-authored and human-authored changes alike.",
       accent=AMBER)
gates = ["Build", "SAST", "Dependency & secrets", "IaC policy", "Compliance evidence", "Human approval", "Deploy"]
gw = 1.68
for i, g in enumerate(gates):
    x = 0.55 + i * (gw + 0.075)
    col = GREEN if g == "Human approval" else AMBER
    sh = box(s, x, 2.05, gw, 0.66, fill=PANEL if g != "Human approval" else C(0x10, 0x2A, 0x24),
             radius=0.12, line=col, lw=1.1)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = I(0.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = g
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(10.5), True, (WHITE if g != "Human approval" else GREEN)

box(s, 0.55, 3.05, 6.00, 2.55, grad=(PANEL, PANEL2), radius=0.07, angle=90, line=C(0x24, 0x33, 0x55), lw=0.75)
txt(s, 0.85, 3.28, 5.4, 0.3, "What the gate enforces", size=14, color=WHITE, bold=True, space=0)
enf = ["Static analysis and dependency risk on every change",
       "No secrets, no unreviewed IaC or policy drift",
       "Scan results attached to the change, not to a spreadsheet",
       "Promotion blocked until a named approver signs off",
       "Rollback path recorded before deploy"]
yy = 3.78
for e in enf:
    box(s, 0.87, yy + 0.07, 0.07, 0.07, fill=AMBER, radius=0.5)
    txt(s, 1.07, yy, 5.2, 0.32, e, size=11, color=MUTED, space=0, line=1.15)
    yy += 0.36

box(s, 6.78, 3.05, 6.00, 2.55, grad=(C(0x10, 0x2A, 0x24), C(0x0E, 0x18, 0x30)), radius=0.07, angle=90,
    line=GREEN, lw=0.9)
txt(s, 7.08, 3.28, 5.4, 0.3, "The person stays in the loop", size=14, color=WHITE, bold=True, space=0)
txt(s, 7.08, 3.66, 5.4, 0.6,
    "The approver sees the ticket, the AI reasoning, the diff and the scan results in one place \u2014 "
    "then approves, rejects, or sends it back.", size=11, color=MUTED, space=0, line=1.2)
chain = ["Ticket", "Session log", "Pull request", "Scan results", "Approval", "Release"]
txt(s, 7.08, 4.40, 5.4, 0.25, "Unbroken audit trail", size=10, color=GREEN, bold=True, caps=True, spacing=1.4, space=0)
for i, c in enumerate(chain):
    x = 7.08 + (i % 3) * 1.85
    y = 4.68 + (i // 3) * 0.48
    chip(s, x, y, 1.62, c, GREEN, h=0.38, size=9.5)
box(s, 0.55, 5.85, 12.23, 0.92, grad=(C(0x2A, 0x1C, 0x08), C(0x0E, 0x18, 0x30)), radius=0.10, angle=0,
    line=AMBER, lw=0.9)
txt(s, 0.85, 6.10, 11.6, 0.45,
    "Nothing AI-authored reaches production without passing the same compliance scans and a named human approval.",
    size=13.5, color=WHITE, bold=True, space=0)
footer(s, "Compliance CI/CD \u00b7 Harness")

# ================= 6. value =================
s = sl()
header(s, "The business case", "Capacity back, risk down, evidence by default.",
       "What changes when qualification, execution and compliance are automated around the same ticket.",
       accent=GREEN)
kpis = [("Ticket deflection", "Repeats resolved by linkage instead of new work", VIOLET),
        ("Cycle time", "Qualified ticket \u2192 approved release", TEAL),
        ("Autonomous fix rate", "Share of tickets resolved without hands-on engineering", CYAN),
        ("Gate pass rate", "AI changes clearing compliance on first submission", AMBER)]
kw = 2.95
for i, (h, b, col) in enumerate(kpis):
    x = 0.55 + i * (kw + 0.115)
    box(s, x, 2.00, kw, 1.70, grad=(PANEL, PANEL2), radius=0.08, angle=90, line=C(0x24, 0x33, 0x55), lw=0.75)
    box(s, x, 2.00, kw, 0.05, fill=col, shape=MSO_SHAPE.RECTANGLE)
    txt(s, x + 0.28, 2.30, kw - 0.5, 0.35, h, size=14, color=WHITE, bold=True, space=0, line=1.1)
    txt(s, x + 0.28, 2.86, kw - 0.5, 0.7, b, size=10.5, color=MUTED, space=0, line=1.2)
    txt(s, x + 0.28, 3.35, kw - 0.5, 0.3, "measured per stage", size=9, color=col, bold=True, space=0)

cols = [("Today", ["Repeat tickets re-investigated from scratch",
                   "Thin tickets bounce between queues for days",
                   "Compliance evidence assembled by hand, late",
                   "Fix quality depends on who picked up the ticket"], C(0x5B, 0x6B, 0x88)),
        ("With this pipeline", ["Repeats linked and closed at intake",
                                "Tickets qualified before capacity is spent",
                                "Scans and approvals captured on every change",
                                "Playbooks make the best method the default"], GREEN)]
for i, (h, items, col) in enumerate(cols):
    x = 0.55 + i * 6.23
    box(s, x, 3.95, 6.00, 2.35, grad=(PANEL, PANEL2) if i == 0 else (C(0x10, 0x2A, 0x24), C(0x0E, 0x18, 0x30)),
        radius=0.07, angle=90, line=C(0x24, 0x33, 0x55) if i == 0 else GREEN, lw=0.8)
    txt(s, x + 0.30, 4.18, 5.4, 0.3, h, size=14.5, color=WHITE if i else MUTED, bold=True, space=0)
    yy = 4.66
    for it in items:
        txt(s, x + 0.30, yy, 0.3, 0.3, "\u2715" if i == 0 else "\u2713", size=11, color=col, bold=True, space=0)
        txt(s, x + 0.62, yy, 5.1, 0.32, it, size=11, color=MUTED, space=0, line=1.15)
        yy += 0.40
txt(s, 0.55, 6.52, 12.23, 0.4,
    "Next step: pick two ticket classes, run the pipeline end to end, and publish the four metrics above.",
    size=12.5, color=WHITE, bold=True, space=0)
footer(s, "Value and next step")

prs.save("AI-Governed-Service-Delivery.pptx")
print("saved")
